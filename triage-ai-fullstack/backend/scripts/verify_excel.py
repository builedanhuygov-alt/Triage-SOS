"""Kiem tra nhanh Excel release."""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
print("sheets:", wb.sheetnames)
required = {"ACCOUNTS", "ROLES", "PERMISSIONS", "ROLE_PERMISSIONS", "ORGANIZATION", "DEPARTMENTS", "BUILDINGS", "DOCTORS", "PATIENT_DEMO", "CAREGIVER_DEMO", "SUMMARY"}
missing = required.difference(wb.sheetnames)
if missing:
      raise SystemExit(f"missing sheets: {sorted(missing)}")
acc = list(wb["ACCOUNTS"].iter_rows(min_row=4, values_only=True))
print("accounts:", len(acc))
users = [r[3] for r in acc]
emps = [r[1] for r in acc]
print("uniq username:", len(set(users)), "uniq emp:", len(set(emps)))
print("no-role:", sum(1 for r in acc if not r[5]), "no-dept:", sum(1 for r in acc if not r[7]))
print("status:", {s: sum(1 for r in acc if r[10] == s) for s in ("ACTIVE", "INACTIVE", "LOCKED", "PENDING")})
print("doctors:", wb["DOCTORS"].max_row - 3, "roles:", wb["ROLES"].max_row - 3,
      "perms:", wb["PERMISSIONS"].max_row - 3, "mappings:", wb["ROLE_PERMISSIONS"].max_row - 3,
      "org rows:", wb["ORGANIZATION"].max_row - 3)
s = {r[0]: r[1] for r in wb["SUMMARY"].iter_rows(min_row=4, values_only=True)}
print("summary total:", s.get("Total Accounts"), "doctors:", s.get("Total Doctors"))
print("patient demos:", wb["PATIENT_DEMO"].max_row - 3, "caregiver demos:", wb["CAREGIVER_DEMO"].max_row - 3)
if wb["PATIENT_DEMO"].max_row - 3 != 3:
      raise SystemExit("PATIENT_DEMO must contain exactly 3 primary demo accounts")
