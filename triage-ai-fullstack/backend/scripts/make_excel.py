"""Tao release/demo-hospital-accounts.xlsx tu JSON export cua lib/hospital.js.

Chay:
    cd triage-ai-fullstack/backend
    node scripts/export-accounts-json.js scripts/accounts.json
    python scripts/make_excel.py scripts/accounts.json ../release/demo-hospital-accounts.xlsx
"""
import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=12)
BANNER_FILL = PatternFill("solid", fgColor="FEF3C7")
DEMO_PASSWORD_HEADER = "Demo Password"


def sheet_setup(ws, title, headers, rows, banner):
    ws.append([banner])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(row=1, column=1)
    c.fill = BANNER_FILL
    c.font = Font(bold=True, size=11, color="92400E")
    c.alignment = Alignment(horizontal="center")
    ws.append([title])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(row=2, column=1).font = TITLE_FONT
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    widths = {}
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            v = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 10), min(48, len(v) + 2))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{ws.max_row}"


def scope_text(u):
    s = u.get("scope") or {}
    parts = [f"LEVEL {s.get('level')}"]
    if s.get("building"):
        parts.append(f"Toa {s['building']}")
    if u.get("departmentName"):
        parts.append(u["departmentName"])
    if u.get("unit"):
        parts.append(u["unit"])
    if s.get("dataAccess"):
        parts.append(s["dataAccess"])
    return " / ".join(parts)


def main(src, dest):
    with open(src, encoding="utf-8") as f:
        d = json.load(f)
    banner = f"{d['banner']} — {d['hospital']}"
    wb = Workbook()

    users = d["users"]
    ws = wb.active
    ws.title = "ACCOUNTS"
    sheet_setup(ws, f"ACCOUNTS — {d['hospital']} ({len(users)} tai khoan demo)",
                ["User ID", "Employee ID", "Full Name", "Username", DEMO_PASSWORD_HEADER,
                 "Role", "Job Title", "Department", "Building", "Floor", "Status",
                 "Email", "Supervisor", "Scope"],
                [[u["userId"], u["employeeId"], u["fullName"], u["username"], d["demoPassword"],
                  u["role"], u["jobTitle"], u.get("departmentName", ""), u.get("building", ""),
                  u.get("floor", ""), u.get("status", ""), u["email"], u.get("supervisor", ""),
                  scope_text(u)] for u in users], banner)

    ws = wb.create_sheet("ROLES")
    sheet_setup(ws, "ROLES",
                ["Role ID", "Role Name", "Description", "Scope"],
                [[r["id"], r["name"], r["description"], f"{r['scope']} (level {r['level']})"] for r in d["roles"]], banner)

    ws = wb.create_sheet("PERMISSIONS")
    sheet_setup(ws, "PERMISSIONS",
                ["Permission ID", "Permission", "Description"],
                [[i + 1, p["id"], p["description"]] for i, p in enumerate(d["permissions"])], banner)

    rows = []
    for r in d["roles"]:
        for p in d["rolePermissions"].get(r["id"], []):
            rows.append([r["id"], p, r["scope"]])
    ws = wb.create_sheet("ROLE_PERMISSIONS")
    sheet_setup(ws, "ROLE_PERMISSIONS", ["Role", "Permission", "Scope"], rows, banner)

    users_by_name = {u["username"]: u["fullName"] for u in users}
    orows = []
    for dep in d["departments"]:
        for fl in dep["floors"]:
            for un in dep["units"]:
                orows.append([dep["building"], fl, dep["name"], un["name"],
                              users_by_name.get(dep.get("head", ""), "")])
    ws = wb.create_sheet("ORGANIZATION")
    sheet_setup(ws, "ORGANIZATION", ["Building", "Floor", "Department", "Unit", "Department Head"], orows, banner)

    ws = wb.create_sheet("DEPARTMENTS")
    sheet_setup(ws, "DEPARTMENTS", ["Department ID", "Code", "Department", "Building", "Floors", "Head", "Units"],
                [[dep["id"], dep["code"], dep["name"], dep["building"], ", ".join(map(str, dep["floors"])), users_by_name.get(dep.get("head", ""), ""), ", ".join(u["name"] for u in dep["units"])] for dep in d["departments"]], banner)

    ws = wb.create_sheet("BUILDINGS")
    sheet_setup(ws, "BUILDINGS", ["Building ID", "Code", "Building", "Floors", "Note"],
                [[b["id"], b["code"], b["name"], ", ".join(map(str, b["floors"])), b.get("note", "")] for b in d["buildings"]], banner)

    ws = wb.create_sheet("DOCTORS")
    sheet_setup(ws, "DOCTORS",
                ["Doctor ID", "Full Name", "Specialty", "Department", "Building", "Position", "Username", "Status"],
                [[x["employeeId"], x["fullName"], x.get("specialty") or "", x.get("departmentName", ""),
                  x.get("building", ""), x.get("jobTitle", ""), x["username"], x.get("status", "")]
                 for x in d["doctors"]], banner)

    ws = wb.create_sheet("PATIENT_DEMO")
    sheet_setup(ws, "PATIENT_DEMO", ["Patient ID", "Medical ID", "Name", "Username", DEMO_PASSWORD_HEADER, "Primary Department", "Doctor", "Building", "Appointment", "Caregiver", "Demo Scenario"],
                [[p["patientId"], p["medicalId"], p["name"], p["username"], p["demoPassword"], p["department"], p["doctor"], p["building"], p["appointment"], p["caregiver"], p["scenario"]] for p in d["patientDemo"]], banner)

    ws = wb.create_sheet("CAREGIVER_DEMO")
    sheet_setup(ws, "CAREGIVER_DEMO", ["Caregiver ID", "Name", "Username", DEMO_PASSWORD_HEADER, "Linked Patient"],
                [[c["caregiverId"], c["name"], c["username"], c["demoPassword"], c["linkedPatient"]] for c in d["caregiverDemo"]], banner)

    ws = wb.create_sheet("PATIENTS")
    sheet_setup(ws, "PATIENTS — ho so demo (DEMO DATA)",
                ["Patient ID", "Full Name", "CCCD", "DOB", "Gender", "Phone", "Blood", "BHYT Code", "BHYT Status", "Allergies", "Conditions"],
                [[p["patientId"], p["fullName"], p.get("cccd", ""), p.get("dob", ""), p.get("gender", ""),
                  p.get("phone", ""), p.get("bloodType", ""), (p.get("bhyt") or {}).get("code", ""),
                  (p.get("bhyt") or {}).get("status", ""), ", ".join(p.get("allergies", [])),
                  ", ".join(p.get("conditions", []))] for p in d.get("patients", [])], banner)

    ws = wb.create_sheet("RESOURCES")
    sheet_setup(ws, "RESOURCES — ton kho demo (DEMO DATA)",
                ["Resource ID", "Name", "Type", "Unit", "Qty", "Min", "Alert At", "Batch", "Expiry", "Building", "Floor", "Department", "Room", "Manager", "Status"],
                [[r["id"], r["name"], r.get("type", ""), r.get("unit", ""), r.get("qty", 0),
                  r.get("minQty", 0), r.get("alertQty", 0), r.get("batch", ""), r.get("expiry", ""),
                  r.get("building", ""), r.get("floor", ""), r.get("department", ""), r.get("room", ""),
                  r.get("manager", ""), r.get("status", "")] for r in d.get("resources", [])], banner)

    ws = wb.create_sheet("SCENARIOS")
    sheet_setup(ws, "SCENARIOS — kich ban demo end-to-end",
                ["#", "Scenario", "Flow", "Accounts"],
                [[1, "Ngoai tru Tim mach", "INTAKE -> ENCOUNTER -> NOTE -> DIAGNOSIS -> LAB ORDER -> RESULT VERIFY -> RELEASE -> RX -> APPROVE -> DISPENSE -> BILL(QR) -> PAY -> FOLLOWUP", "emergency.head.demo / cardio doctor / pharmacy / finance"],
                 [2, "Cap cuu SOS", "SOS_REQUESTED -> ACKNOWLEDGED -> DISPATCHING -> EN_ROUTE -> ARRIVED -> CLOSED (DEMO/SIMULATED)", "Emergency Desk (emergency.head.demo)"],
                 [3, "Xuat kho vaccine 2 cap", "REQUEST (Truong khoa) -> PROCESS (Kho) -> APPROVE (Director) -> qty tru + alert LOW/CRITICAL", "emergency.head.demo / pharmacy / director.demo"],
                 [4, "Thuoc theo don", "PRESCRIPTION -> SCHEDULE -> REMINDER -> TAKEN (xac nhan) / MISSED (sweep)", "patient@hospital.demo"],
                 [5, "BHYT thu cong", "VERIFY -> MANUAL_REVIEW -> manual VERIFIED (ADAPTER READY, DEMO/SIMULATED)", "finance staff"],
                 [6, "Anomaly review", "3 login fail -> ANOMALY_DETECTED (OPEN) -> REVIEW -> RESOLVED (khong tu khoa TK)", "director.demo"],
                 [7, "Van thu 2 cap", "CREATE -> RECEIVE -> PROCESS -> PENDING_APPROVAL -> APPROVE -> DONE", "admin staff / planning.manager.demo / director.demo"]], banner)

    s = d["summary"]
    ws = wb.create_sheet("SUMMARY")
    sheet_setup(ws, "SUMMARY — thong ke tai khoan demo", ["Metric", "Value"],
                [["Hospital", d["hospital"]], ["Total Accounts", s["totalAccounts"]],
                 ["Active", s["active"]], ["Inactive", s["inactive"]],
                 ["Locked", s["locked"]], ["Pending", s["pending"]],
                 ["Total Doctors", s["doctors"]], ["Total Nurses", s["nurses"]],
                 ["Total Technicians", s["technicians"]], ["Total Pharmacists", s["pharmacists"]],
                 ["Total Admin/Support", s["admin"]], ["Total IT", s["it"]],
                 ["Total Managers", s["managers"]], ["Total Finance", s["finance"]],
                 ["Total HR", s["hr"]], ["Total Departments", s["departments"]],
                 ["Total Buildings", s["buildings"]], ["Total Roles", s["roles"]],
                 ["Total Permissions", s["permissions"]], ["Demo Patients", len(d["patientDemo"])],
                 ["Demo Caregivers", len(d["caregiverDemo"])]], banner)

    wb.save(dest)
    print(f"saved {dest}: {len(users)} accounts, {len(d['roles'])} roles, "
          f"{len(d['permissions'])} permissions, {len(rows)} mappings, "
          f"{len(d['departments'])} departments, {len(d['doctors'])} doctors")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
